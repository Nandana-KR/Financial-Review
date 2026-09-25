"""
Ingest layer: parse a bank transaction file (.xlsx or .csv) into structured rows.

Bank exports vary wildly in column naming, so this uses fuzzy header matching and
supports both the "single signed amount" layout and the "separate debit/credit"
layout. Everything is normalized to a single signed `amount`:
    positive = money in (inflow / credit)
    negative = money out (outflow / debit)

This layer does NO categorization. It only produces clean structured records.
"""

from __future__ import annotations
import csv
import io
import json
import re
from dataclasses import dataclass, asdict
from datetime import datetime, date
from typing import Any, Optional

from dateutil import parser as dateparser
from openpyxl import load_workbook


@dataclass
class ParsedTxn:
    txn_date: str          # ISO YYYY-MM-DD
    period: str            # YYYY-MM
    description: str
    counterparty: Optional[str]
    amount: float          # signed, positive=inflow
    raw_amount: float
    account: Optional[str]
    raw_json: str


# Header aliases -> canonical field. Lowercased, stripped for matching.
HEADER_ALIASES = {
    "date": {"date", "transaction date", "txn date", "posted date", "post date", "value date"},
    "description": {"description", "details", "memo", "narrative", "particulars", "transaction", "name"},
    "counterparty": {"counterparty", "merchant", "payee", "vendor", "payer", "to/from"},
    "amount": {"amount", "value", "transaction amount", "amt"},
    "debit": {"debit", "withdrawal", "withdrawals", "money out", "paid out", "outflow"},
    "credit": {"credit", "deposit", "deposits", "money in", "paid in", "inflow"},
    "account": {"account", "account name", "account number", "acct"},
    "type": {"type", "transaction type", "dr/cr", "direction"},
}


def _norm_header(h: Any) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower())


def _match_columns(headers: list[str]) -> dict[str, int]:
    """Map canonical field name -> column index based on header aliases."""
    mapping: dict[str, int] = {}
    for idx, h in enumerate(headers):
        nh = _norm_header(h)
        for field, aliases in HEADER_ALIASES.items():
            if nh in aliases and field not in mapping:
                mapping[field] = idx
    return mapping


def _parse_amount(val: Any) -> Optional[float]:
    if val is None or val == "":
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if s == "":
        return None
    negative = False
    # Parentheses accounting notation e.g. (1,234.56) means negative.
    if s.startswith("(") and s.endswith(")"):
        negative = True
        s = s[1:-1]
    s = s.replace("$", "").replace(",", "").replace(" ", "")
    if s.startswith("-"):
        negative = True
        s = s[1:]
    try:
        num = float(s)
    except ValueError:
        return None
    return -num if negative else num


def _parse_date(val: Any) -> Optional[date]:
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    try:
        return dateparser.parse(str(val), dayfirst=False).date()
    except (ValueError, OverflowError):
        try:
            return dateparser.parse(str(val), dayfirst=True).date()
        except (ValueError, OverflowError):
            return None


def _rows_from_xlsx(content: bytes) -> list[list[Any]]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def _rows_from_csv(content: bytes) -> list[list[Any]]:
    text = content.decode("utf-8-sig", errors="replace")
    return [row for row in csv.reader(io.StringIO(text))]


def _find_header_row(rows: list[list[Any]]) -> int:
    """Find the row index that looks most like a header (matches known aliases)."""
    best_idx, best_score = 0, -1
    for i, row in enumerate(rows[:15]):  # headers are near the top
        headers = [_norm_header(c) for c in row]
        score = 0
        for aliases in HEADER_ALIASES.values():
            if any(h in aliases for h in headers):
                score += 1
        if score > best_score:
            best_idx, best_score = i, score
    return best_idx


def parse_file(filename: str, content: bytes) -> list[ParsedTxn]:
    """Parse an uploaded file into normalized transactions."""
    lower = filename.lower()
    if lower.endswith(".csv"):
        rows = _rows_from_csv(content)
    elif lower.endswith((".xlsx", ".xlsm")):
        rows = _rows_from_xlsx(content)
    else:
        # Try xlsx then csv as best effort.
        try:
            rows = _rows_from_xlsx(content)
        except Exception:
            rows = _rows_from_csv(content)

    rows = [r for r in rows if r and any(c is not None and str(c).strip() != "" for c in r)]
    if not rows:
        return []

    header_idx = _find_header_row(rows)
    headers = [str(c) if c is not None else "" for c in rows[header_idx]]
    mapping = _match_columns(headers)
    data_rows = rows[header_idx + 1:]

    parsed: list[ParsedTxn] = []
    for row in data_rows:
        record = _build_record(row, headers, mapping)
        if record is not None:
            parsed.append(record)
    return parsed


def _build_record(row: list[Any], headers: list[str], mapping: dict[str, int]) -> Optional[ParsedTxn]:
    def get(field: str) -> Any:
        idx = mapping.get(field)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    d = _parse_date(get("date"))
    if d is None:
        return None  # rows without a valid date are not transactions

    # Resolve signed amount.
    signed: Optional[float] = None
    raw_amount: Optional[float] = None

    if "amount" in mapping:
        amt = _parse_amount(get("amount"))
        raw_amount = amt
        signed = amt
        # If a type/direction column exists and amount is unsigned, apply sign.
        tval = get("type")
        if signed is not None and tval is not None:
            t = str(tval).strip().lower()
            if t in {"debit", "dr", "withdrawal", "out"} and signed > 0:
                signed = -signed
            elif t in {"credit", "cr", "deposit", "in"} and signed < 0:
                signed = abs(signed)

    if signed is None and ("debit" in mapping or "credit" in mapping):
        debit = _parse_amount(get("debit")) or 0.0
        credit = _parse_amount(get("credit")) or 0.0
        # Debit = money out (negative), credit = money in (positive).
        signed = credit - abs(debit)
        raw_amount = signed

    if signed is None:
        return None  # no interpretable amount

    description = str(get("description") or "").strip()
    counterparty_raw = get("counterparty")
    counterparty = str(counterparty_raw).strip() if counterparty_raw else None
    if not description:
        description = counterparty or "(no description)"

    account_raw = get("account")
    account = str(account_raw).strip() if account_raw else None

    raw_json = json.dumps(
        {headers[i] if i < len(headers) else f"col{i}": _jsonable(row[i])
         for i in range(len(row))},
        default=str,
    )

    return ParsedTxn(
        txn_date=d.isoformat(),
        period=f"{d.year:04d}-{d.month:02d}",
        description=description,
        counterparty=counterparty,
        amount=round(float(signed), 2),
        raw_amount=round(float(raw_amount if raw_amount is not None else signed), 2),
        account=account,
        raw_json=raw_json,
    )


def _jsonable(v: Any) -> Any:
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return v


def to_dicts(txns: list[ParsedTxn]) -> list[dict]:
    return [asdict(t) for t in txns]
